# -*- coding: utf-8 -*-
"""Genera la monografía UniGuajira News en formato .docx (ICONTEC NTC 1486)."""

import re
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = "/home/miguel/uniguajira-news/monografia/contenido"
OUT = "/home/miguel/uniguajira-news/monografia/Monografia_UniGuajira_News_2026.docx"

BODY_FILES = [
    "01_introduccion.md",
    "02_capitulo1.md",
    "03_capitulo2.md",
    "04_capitulo3.md",
    "05_capitulo4.md",
    "06_capitulo5.md",
    "07_conclusiones.md",
    "08_referencias_anexos.md",
]


def read_file(name):
    with open(f"{BASE}/{name}", encoding="utf-8") as f:
        return f.read()


def set_margins(section, top, bottom, left, right):
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)


def field_parts(paragraph, instr):
    r1 = paragraph.add_run()
    fbeg = OxmlElement("w:fldChar"); fbeg.set(qn("w:fldCharType"), "begin")
    r1._r.append(fbeg)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    fsep = OxmlElement("w:fldChar"); fsep.set(qn("w:fldCharType"), "separate")
    r3._r.append(fsep)
    r4 = paragraph.add_run()
    fend = OxmlElement("w:fldChar"); fend.set(qn("w:fldCharType"), "end")
    r4._r.append(fend)
    return [r1, r2, r3, r4]


def add_page_number(section, start_at=1):
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), str(start_at))
    section._sectPr.append(pg)
    section.footer.is_linked_to_previous = False
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    runs = field_parts(para, "PAGE")
    for r in runs:
        r.font.name = "Arial"
        r.font.size = Pt(12)


def set_font(run, size=12, bold=False, italic=False, name="Arial", color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def para(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, bold=False,
         italic=False, before=0, after=6, line=1.5, name="Arial"):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if text:
        run = p.add_run(text)
        set_font(run, size=size, bold=bold, italic=italic, name=name)
    return p


def heading(doc, text, level=1, page_break=False):
    p = doc.add_paragraph()
    text = text.strip()
    pf = p.paragraph_format
    if page_break:
        pf.page_break_before = True
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.space_after = Pt(14)
        pf.line_spacing = 1.5
        run = p.add_run(text.upper().rstrip("."))
        set_font(run, size=12, bold=True)
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(12)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        run = p.add_run(text.upper().rstrip("."))
        set_font(run, size=12, bold=True)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.line_spacing = 1.5
        run = p.add_run(text)
        set_font(run, size=12, bold=True)
    return p


def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def add_table(doc, rows):
    rows = [r for r in rows if r]
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    for i, row in enumerate(rows):
        for j in range(cols):
            txt = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            p.text = ""
            run = p.add_run(txt)
            set_font(run, size=10, bold=(i == 0))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.1
    para(doc, "", after=4)


def render_md(doc, md):
    lines = md.splitlines()
    i = 0
    n = len(lines)
    in_code = False
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue

        if in_code:
            para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, size=10, name="Courier New")
            i += 1
            continue

        if stripped.startswith("###"):
            heading(doc, stripped.lstrip("#").strip(), level=3)
            i += 1
            continue

        if stripped.startswith("##"):
            text = stripped.lstrip("#").strip()
            if text.startswith(("Tabla ", "Figura ")):
                para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
            else:
                heading(doc, text, level=2)
            i += 1
            continue

        if stripped.startswith("#"):
            heading(doc, stripped.lstrip("#").strip(), level=1, page_break=True)
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", c) for c in row):
                    i += 1
                    continue
                rows.append(row)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            run = p.add_run(stripped[2:])
            set_font(run, size=12)
            i += 1
            continue

        if stripped.startswith(("PALABRAS CLAVE", "KEYWORDS")):
            key, _, rest = stripped.partition(":")
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            r1 = p.add_run(key + ":")
            r2 = p.add_run(rest)
            set_font(r1, size=12, bold=True)
            set_font(r2, size=12)
            i += 1
            continue

        if stripped.startswith("Fuente:"):
            para(doc, stripped, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, italic=True, after=12)
            i += 1
            continue

        if stripped.startswith(("Tabla ", "Figura ")) and len(stripped) < 120:
            para(doc, stripped, align=WD_ALIGN_PARAGRAPH.CENTER, size=11, after=4)
            i += 1
            continue

        para(doc, stripped)
        i += 1


def build_cover(doc):
    for _ in range(2):
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "UNIVERSIDAD DE LA GUAJIRA", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True)
    para(doc, "SEDE MAICAO", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "FACULTAD DE INGENIERÍA — PROGRAMA DE INGENIERÍA DE SISTEMAS",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "DESARROLLO DE UNA APLICACIÓN MÓVIL PARA LA CENTRALIZACIÓN",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, line=1.3)
    para(doc, "Y DIFUSIÓN DE NOTICIAS INSTITUCIONALES DE LA UNIVERSIDAD",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, line=1.3)
    para(doc, "DE LA GUAJIRA", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, line=1.3)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "MIGUEL ÁNGEL URECHE BRAVO", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "JHOSTIN PABÓN RAMÍREZ", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Monografía con Desarrollo Tecnológico presentada como requisito",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=12, line=1.3)
    para(doc, "para optar al título de Ingeniero de Sistemas",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=12, line=1.3)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Línea de investigación: Desarrollo de Software, Sistemas de Información",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=11, line=1.3)
    para(doc, "y Transformación Digital Regional", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, line=1.3)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Director:", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    para(doc, "FARITH PEREZ SAEZ", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "MAICAO — LA GUAJIRA", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "2026", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)


def build_acceptance(doc):
    doc.add_page_break()
    for _ in range(2):
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "UNIVERSIDAD DE LA GUAJIRA", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True)
    para(doc, "FACULTAD DE INGENIERÍA — PROGRAMA DE INGENIERÍA DE SISTEMAS",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "PÁGINA DE ACEPTACIÓN", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, bold=True)
    para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "La sustentación del trabajo de grado titulado DESARROLLO DE UNA APLICACIÓN MÓVIL "
              "PARA LA CENTRALIZACIÓN Y DIFUSIÓN DE NOTICIAS INSTITUCIONALES DE LA UNIVERSIDAD DE "
              "LA GUAJIRA, presentado por Miguel Ángel Ureche Bravo y Jhostin Pabón Ramírez, fue "
              "evaluada y aprobada por el jurado designado.",
         align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, line=1.5)
    for _ in range(5):
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(3):
        para(doc, "_________________________________________",
             align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
        para(doc, "Jurado evaluador", align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)


def build_dedication(doc):
    doc.add_page_break()
    for _ in range(6):
        para(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "A nuestras familias, que sostuvieron este proceso desde el primer semestre "
              "y festejaron cada avance como propio.",
         align=WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, line=2.0)


def build_acknowledgments(doc):
    doc.add_page_break()
    heading(doc, "AGRADECIMIENTOS", level=1)
    para(doc, "Esta monografía no habría llegado a puerto sin el acompañamiento de varias "
              "personas. En primer lugar, agradecemos a nuestro director, Farith Perez Saez, "
              "por guiar el trabajo con paciencia y por exigir rigor sin "
              "perder de vista el contexto real en el que viven los estudiantes de la Sede "
              "Maicao.", size=12)
    para(doc, "Agradecemos también al personal de sistemas de la universidad, que nos abrió "
              "las puertas del portal y respondió con franqueza a nuestras entrevistas, y a "
              "los treinta estudiantes y docentes que participaron en la prueba piloto: sus "
              "respuestas son la razón de ser de este proyecto.", size=12)
    para(doc, "Por último, a nuestras familias y amigos, que entendieron las ausencias y las "
              "madrugadas de los últimos meses, y a la Universidad de La Guajira, que nos "
              "formó como ingenieros y nos permitió devolverle un producto útil.", size=12)


def build_toc_field(doc):
    p = doc.add_paragraph()
    runs = field_parts(p, 'TOC \\o "1-3" \\h \\z \\u')
    ph = p.add_run("  [En Word: clic derecho aquí → Actualizar campo. ]")
    set_font(ph, size=10, italic=True)
    runs.append(ph)


def main():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")

    sec0 = doc.sections[0]
    set_margins(sec0, top=4, bottom=3, left=4, right=2)
    sec0.footer.is_linked_to_previous = False
    sec0.footer.paragraphs[0].text = ""

    build_cover(doc)
    build_acceptance(doc)
    build_dedication(doc)
    build_acknowledgments(doc)

    render_md(doc, read_file("00_preliminares.md"))

    heading(doc, "CONTENIDO", level=1)
    build_toc_field(doc)

    sec1 = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_margins(sec1, top=3, bottom=3, left=4, right=2)
    add_page_number(sec1, start_at=1)

    for name in BODY_FILES:
        render_md(doc, read_file(name))

    doc.save(OUT)
    print(f"OK: {OUT}")


if __name__ == "__main__":
    main()